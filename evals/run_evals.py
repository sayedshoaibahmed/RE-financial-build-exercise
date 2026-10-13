"""run_evals.py: the Task 3 eval suite. Offline and deterministic: no network, no model.

    python evals/run_evals.py                    run every case and every commentary draft, diff against the baseline
    python evals/run_evals.py G03 X02            run only the named cases (commentary grading always runs)
    python evals/run_evals.py --update-baseline  print the diff, then rewrite evals/baseline.json (full run only)
    python evals/run_evals.py --help             all options (--cases-dir, --results, --baseline, --no-baseline)

EVALS_SKILL_DIR=<folder> runs the suite against another copy of the skill (evals/run_mutants.py uses it).

Part 1 (golden cases, evals/cases/*.json): each run drives the real CLI (trend_cli.py run / render) on ABC or on a
transformed copy of the ABC workbook, then checks status, exit code, pinned facts, n/a reasons, flags, blocked reports
and file effects. Every pinned number is also recomputed by the independent oracle (evals/oracle_facts.py over
tests/oracle.py); a case whose literal disagrees with the oracle fails, so expectations never come from the package.

Part 2 (commentary grading): every good draft in tests/fixtures/commentary/ must pass `render`; every bad draft in
evals/bad_drafts/g<N>_*.md must fail `render` with exactly the code G<N> and nothing else.

Part 3 (regression baseline): the results are compared with evals/baseline.json (every row, every check name with
its result, and a hash of each case file and draft). Any difference is printed and fails the run.

Prints each row as it finishes, a PASS/FAIL table (a short diff under each failure), the baseline diff and the
elapsed time; writes evals/results/latest.json. Exit 0: every row passed and the baseline matches. 1: a row failed.
2: bad arguments. 5: no usable Python. 6: every row passed but the results differ from the baseline.
"""
import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SKILL_DIR = Path(os.environ.get("EVALS_SKILL_DIR") or ROOT / ".claude" / "skills" / "monthly-trend-analysis").resolve()
SKILL_SCRIPTS = SKILL_DIR / "scripts"
CLI = SKILL_SCRIPTS / "trend_cli.py"


def _ensure_interpreter():
    """Re-run under the project's Python (the one ensure_python.py picks) if this one lacks openpyxl / yaml."""
    try:
        import openpyxl  # noqa: F401
        import yaml  # noqa: F401
    except ImportError:
        r = subprocess.run([sys.executable, str(SKILL_SCRIPTS / "ensure_python.py"), "--which"], capture_output=True, text=True)
        exe = r.stdout.strip()
        if r.returncode != 0 or not exe or Path(exe).resolve() == Path(sys.executable).resolve():
            print("error: no Python with openpyxl and PyYAML found.\n" + r.stderr, file=sys.stderr)
            sys.exit(5)
        sys.exit(subprocess.run([exe, str(Path(__file__).resolve()), *sys.argv[1:]]).returncode)


_ensure_interpreter()

import openpyxl  # noqa: E402
import yaml  # noqa: E402

sys.path[:0] = [str(ROOT), str(SKILL_SCRIPTS)]        # tests.fixtures.mini (the mini company builder) imports trend
from evals.oracle_facts import Oracle, text2  # noqa: E402

CASES = ROOT / "evals" / "cases"
BASELINE = ROOT / "evals" / "baseline.json"
BAD = ROOT / "evals" / "bad_drafts"
GOOD_DIR = ROOT / "tests" / "fixtures" / "commentary"
RESULTS = ROOT / "evals" / "results" / "latest.json"
ABC_DIR = ROOT / "companies" / "abc_inc"
SHEET, HEADER_ROW = "Income Statement", 5
GUARDED = [ABC_DIR / "company.yaml", ABC_DIR / "mapping.yaml", ROOT / "tests" / "golden" / "abc_flags.expected.json"]
GOLDEN_FIELDS = ("id", "type", "severity", "period", "subject", "observed", "basis")


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def cli(*args):
    return subprocess.run([sys.executable, str(CLI), *map(str, args)], cwd=ROOT, capture_output=True, text=True)


class Checks:
    """Collects (name, ok, detail) for one case."""

    def __init__(self):
        self.items = []

    def add(self, name, ok, detail=""):
        self.items.append({"check": name, "ok": bool(ok), "detail": "" if ok else detail})
        return ok

    def eq(self, name, actual, expected):
        return self.add(name, actual == expected, f"expected {expected!r}, got {actual!r}")

    @property
    def ok(self):
        return all(i["ok"] for i in self.items)


# ---------------------------------------------------------------- inputs: the real company or a transformed copy

def _label_row(ws, label):
    hits = [r for r in range(1, ws.max_row + 1)
            for c in (1, 2) if isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip() == label]
    if len(hits) != 1:
        raise ValueError(f"transform: expected one row labelled {label!r}, found {len(hits)}")
    return hits[0]


def _period_col(ws, period):
    from datetime import datetime
    for c in range(1, ws.max_column + 1):
        v = ws.cell(HEADER_ROW, c).value
        if isinstance(v, str):
            try:
                if datetime.strptime(v.strip(), "%b %Y").strftime("%Y-%m") == period:
                    return c
            except ValueError:
                pass
    raise ValueError(f"transform: no column for {period}")


def _period_cols(ws):
    from datetime import datetime
    out = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(HEADER_ROW, c).value
        if isinstance(v, str):
            try:
                datetime.strptime(v.strip(), "%b %Y")
                out.append(c)
            except ValueError:
                pass
    return out


def transform_workbook(src, dst, t):
    """Documented, reproducible edits to a copy of the workbook. Values are kept (formulas become their cached values)."""
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb[SHEET]
    if "relabel" in t:
        r = _label_row(ws, t["relabel"]["from"])
        col = 1 if isinstance(ws.cell(r, 1).value, str) and ws.cell(r, 1).value.strip() == t["relabel"]["from"] else 2
        ws.cell(r, col).value = t["relabel"]["to"]
    if "add_to_cell" in t:
        a = t["add_to_cell"]
        cell = ws.cell(_label_row(ws, a["label"]), _period_col(ws, a["period"]))
        cell.value = float(Decimal(str(cell.value)) + Decimal(a["amount"]))
    if "insert_row" in t:
        a = t["insert_row"]                                    # amount: same value every month; copy_from: another row's values
        cols = _period_cols(ws)
        values = ([ws.cell(_label_row(ws, a["copy_from"]), c).value for c in cols] if "copy_from" in a
                  else [float(a["amount"])] * len(cols))
        r = _label_row(ws, a["before"])
        ws.insert_rows(r)
        ws.cell(r, 2).value = a["label"]
        for c, v in zip(cols, values):
            ws.cell(r, c).value = v
    if "delete_period" in t:
        ws.delete_cols(_period_col(ws, t["delete_period"]))
    wb.save(dst)


def prepare(inp, tmp):
    """-> (company_dir, workbook path, mapping path). Untransformed ABC runs from companies/abc_inc itself;
    "fixture:<name>" builds a synthetic company (evals/fixture_companies.py) in the temp folder."""
    t = inp.get("transform") or {}
    if inp["company"].startswith("fixture:"):
        from evals.fixture_companies import BUILDERS
        if t:
            raise ValueError("transforms apply to abc_inc only")
        cdir = BUILDERS[inp["company"].split(":", 1)[1]](tmp / "fixture")
        doc = yaml.safe_load((cdir / "company.yaml").read_text(encoding="utf-8"))
        return cdir, Path(doc["source"]["path"]), cdir / "mapping.yaml"
    if inp["company"] != "abc_inc":
        raise ValueError("only abc_inc and fixture:<name> inputs are supported")
    src_doc = yaml.safe_load((ABC_DIR / "company.yaml").read_text(encoding="utf-8"))
    src_wb = ROOT / src_doc["source"]["path"]
    if not t:
        return ABC_DIR, src_wb, ABC_DIR / "mapping.yaml"
    cdir = tmp / "companies" / "abc_inc"
    cdir.mkdir(parents=True)
    wb = tmp / "abc_transformed.xlsx"
    transform_workbook(src_wb, wb, t)
    src_doc["source"]["path"] = str(wb)
    (cdir / "company.yaml").write_text(yaml.safe_dump(src_doc, sort_keys=False), encoding="utf-8")
    mapping = (ABC_DIR / "mapping.yaml").read_text(encoding="utf-8")
    for e in t.get("mapping_add", []):
        mapping = mapping.rstrip("\n") + "\n" + yaml.safe_dump([e], sort_keys=False, default_flow_style=None).rstrip("\n") + "\n"
    (cdir / "mapping.yaml").write_text(mapping, encoding="utf-8")
    return cdir, wb, cdir / "mapping.yaml"


# ---------------------------------------------------------------- one run of one case

def _run_folder(out_base, company_id):
    folders = [d for d in (out_base / company_id).iterdir() if d.is_dir()] if (out_base / company_id).is_dir() else []
    return folders[0] if len(folders) == 1 else None


def _oracle_value(spec, oracles, transform):
    where, _, rest = spec.partition(":")
    if where == "transform":
        return Decimal(transform["add_to_cell"][rest])
    o = oracles[where]
    if rest.startswith("leaf:"):
        _, label, period = rest.split(":")
        return o.leaf(label, period)
    return o.value(rest)


def check_ok_run(chk, exp, run_dir, oracle):
    doc = json.loads((run_dir / "metrics.json").read_text(encoding="utf-8"))
    facts, na = doc["facts"], doc["na"]
    flags = doc["flags"]["flags"]
    if "as_of" in exp:
        chk.eq("as_of", doc["run"]["as_of"], exp["as_of"])
    if "max_period" in exp:
        late = sorted({f.rsplit(".", 1)[1] for f in facts if f.rsplit(".", 1)[1] > exp["max_period"]})
        chk.add("no fact after " + exp["max_period"], not late and max(doc["run"]["periods"]) == exp["max_period"],
                f"facts exist for later periods {late[:5]}; periods end {max(doc['run']['periods'])}")
    for fid, want in exp.get("facts", {}).items():
        got = text2(facts[fid]) if facts.get(fid) is not None else None
        ref = text2(oracle.value(fid))
        if not chk.add(f"oracle agrees {fid}", ref == want, f"case file says {want}, oracle says {ref}: fix the case file"):
            continue
        chk.add(f"fact {fid}", got == want, f"expected {want}, got {got}" + (" (missing)" if fid not in facts else ""))
    for fid, reason in exp.get("na", {}).items():
        ref = oracle.value(fid)
        chk.add(f"oracle agrees {fid} is n/a", ref is None, f"oracle has a value {text2(ref)}: fix the case file")
        present = fid in facts and facts[fid] is None
        why = na.get(fid, "")
        chk.add(f"n/a {fid}" + (f" ({reason})" if reason else ""), present and (reason is None or why.startswith(reason)),
                f"expected n/a{' ' + reason if reason else ''}, got " + (repr(facts[fid]) if fid in facts else "missing")
                + (f" / reason {why!r}" if why else ""))
    for fid in exp.get("no_value", []):
        chk.add(f"no value {fid}", facts.get(fid) is None, f"expected no value, got {facts.get(fid)}")

    def match(f, want):
        return all(f.get(k) == v for k, v in want.items())
    for want in exp.get("flags_required", []):
        chk.add(f"flag present {want}", any(match(f, want) for f in flags), "no such flag")
    for want in exp.get("flags_forbidden", []):
        hits = [f["id"] for f in flags if match(f, want)]
        chk.add(f"flag absent {want}", not hits, f"found {hits[:5]}")
    for typ, n in exp.get("flag_counts", {}).items():
        chk.eq(f"flag count {typ}", sum(f["type"] == typ for f in flags), n)
    if exp.get("digest_contains"):
        digest = (run_dir / "digest.json").read_text(encoding="utf-8")
        for text in exp["digest_contains"]:
            chk.add(f"digest.json carries {text[:40]!r}...", json.dumps(text, ensure_ascii=False)[1:-1] in digest,
                    "text not found in digest.json")
    if "flags_golden" in exp:
        golden = json.loads((ROOT / exp["flags_golden"]).read_text(encoding="utf-8"))["flags"]
        old = {f["id"]: f for f in golden}
        new = {f["id"]: {k: f[k] for k in GOLDEN_FIELDS} for f in flags}
        diff = ([f"ADDED   {i} {new[i]['type']} {new[i]['period']} {new[i]['subject']}" for i in sorted(set(new) - set(old))]
                + [f"REMOVED {i} {old[i]['type']} {old[i]['period']} {old[i]['subject']}" for i in sorted(set(old) - set(new))]
                + [f"CHANGED {i} {k}: {old[i][k]} -> {new[i][k]}" for i in sorted(set(old) & set(new))
                   for k in GOLDEN_FIELDS if old[i][k] != new[i][k]])
        chk.add(f"flags equal {exp['flags_golden']} ({len(golden)} flags)", not diff,
                f"{len(diff)} difference(s): " + "; ".join(diff[:6]) + (" ..." if len(diff) > 6 else ""))


def check_blocked_run(chk, exp, run_dir, oracles, transform, mapping_path, mapping_before):
    rep = json.loads((run_dir / "blocked_report.json").read_text(encoding="utf-8"))
    chk.eq("status BLOCKED", rep.get("status"), "BLOCKED")
    if "blocked_reason" in exp:
        chk.eq("blocked reason", rep.get("reason"), exp["blocked_reason"])
    failed = rep.get("failed_controls", [])
    for want in exp.get("failed_controls", []):
        for field, spec in (want.get("oracle") or {}).items():
            ref = text2(_oracle_value(spec, oracles, transform))
            chk.add(f"oracle agrees {want['control_id']} {want['period']} {field}", ref == want[field],
                    f"case file says {want[field]}, oracle ({spec}) says {ref}: fix the case file")
        hit = [f for f in failed if f["control_id"] == want["control_id"] and f["period"] == want["period"]]
        label = f"control {want['control_id']} fails {want['period']}"
        if not chk.add(label, len(hit) == 1, f"not among the {len(failed)} failed controls"):
            continue
        for field in ("expected", "actual", "difference"):
            if field in want:
                chk.eq(f"{label}: {field}", text2(hit[0][field]), want[field])
    for cid in exp.get("control_fails_every_period", []):
        periods = set(oracles["transformed"].o["periods"])
        got = {f["period"] for f in failed if f["control_id"] == cid}
        chk.add(f"control {cid} fails in every month ({len(periods)})", got == periods,
                f"fails in {len(got)} of {len(periods)}; missing {sorted(periods - got)[:5]}")
    if "failed_periods_only" in exp:
        other = sorted({f["period"] for f in failed} - set(exp["failed_periods_only"]))
        chk.add(f"failures only in {exp['failed_periods_only']}", not other, f"also fails in {other[:5]}")
    if exp.get("failed_controls_have_source_ref"):
        chk.add("every failed control has a cell reference", failed and all(f.get("source_ref") for f in failed),
                "a failed control has no source_ref")
    for acct in exp.get("unmapped_accounts_include", []):
        p = run_dir / "unmapped_accounts.yaml"
        chk.add(f"unmapped_accounts.yaml lists {acct!r}", p.is_file() and acct in p.read_text(encoding="utf-8"),
                "file missing" if not p.is_file() else "account not listed")
    if exp.get("mapping_unchanged"):
        chk.eq("mapping.yaml byte-identical after the run", sha(mapping_path), mapping_before)


def do_render(chk, step, run_dir, n):
    tag = f"render {n}"
    draft = (ROOT / step["draft"]).read_text(encoding="utf-8")
    for old, new in step.get("edits", []):
        if not chk.add(f"{tag}: edit applies once", draft.count(old) == 1, f"{old!r} found {draft.count(old)} times"):
            return
        draft = draft.replace(old, new)
    (run_dir / "commentary_draft.md").write_text(draft, encoding="utf-8", newline="\n")
    if step.get("stale_commentary"):
        (run_dir / "commentary.md").write_text("stale commentary from an earlier render\n", encoding="utf-8")
    if "tamper_digest" in step:
        d = run_dir / "digest.json"
        raw, t = d.read_bytes(), step["tamper_digest"]
        if not chk.add(f"{tag}: tamper target found in digest", t["find"].encode() in raw, f"{t['find']!r} not in digest.json"):
            return
        d.write_bytes(raw.replace(t["find"].encode(), t["replace"].encode(), 1))
    r = cli("render", "--run-dir", run_dir)
    e = step["expect"]
    chk.add(f"{tag}: exit {e['exit_code']}", r.returncode == e["exit_code"],
            f"exit {r.returncode}; stderr: {r.stderr.strip()[-300:]}")
    report = run_dir / "grounding_report.json"
    codes = sorted({f["code"] for f in json.loads(report.read_text(encoding="utf-8"))["failures"]}) if report.is_file() else None
    chk.eq(f"{tag}: failure codes {e['codes']}", codes, sorted(e["codes"]))
    chk.eq(f"{tag}: commentary.md {'written' if e['commentary_written'] else 'absent'}",
           (run_dir / "commentary.md").is_file(), e["commentary_written"])


def one_run(chk, run, tmp):
    cdir, wb, mapping = prepare(run["input"], tmp)
    mapping_before = sha(mapping)
    out = tmp / "out"
    r = cli("run", "--company-dir", cdir, "--out", out, *run["input"].get("args", []))
    exp = run["expect"]
    if not chk.add(f"run exit {exp['exit_code']}", r.returncode == exp["exit_code"],
                   f"exit {r.returncode}; stderr: {r.stderr.strip()[-300:]}"):
        return
    run_dir = _run_folder(out, cdir.name)
    if not chk.add("one run folder", run_dir is not None, "no single run folder under --out"):
        return
    if "run_folder_files" in exp:
        chk.eq(f"run folder holds only {exp['run_folder_files']}",
               sorted(p.name for p in run_dir.iterdir()), sorted(exp["run_folder_files"]))
    oracles = {"transformed": Oracle(wb, cdir / "company.yaml", mapping)}
    if run["input"]["company"] == "abc_inc":
        ref_doc = yaml.safe_load((ABC_DIR / "company.yaml").read_text(encoding="utf-8"))
        oracles["original"] = Oracle(ROOT / ref_doc["source"]["path"], ABC_DIR / "company.yaml", ABC_DIR / "mapping.yaml")
    else:
        oracles["original"] = oracles["transformed"]
    if exp["status"] == "OK":
        check_ok_run(chk, exp, run_dir, oracles["transformed"])
    else:
        check_blocked_run(chk, exp, run_dir, oracles, run["input"].get("transform") or {}, mapping, mapping_before)
    for n, step in enumerate(run.get("render", []), start=1):
        do_render(chk, step, run_dir, n)


def run_case(case):
    chk = Checks()
    for run in case["runs"]:
        start = len(chk.items)
        with tempfile.TemporaryDirectory(prefix="evals_") as td:
            try:
                one_run(chk, run, Path(td))
            except Exception as e:                                         # a broken case is a failure, not a crash
                chk.add("ran without error", False, f"{type(e).__name__}: {e}")
        if len(case["runs"]) > 1:
            for item in chk.items[start:]:
                item["check"] = f"[{run['name']}] {item['check']}"
    return chk


# ---------------------------------------------------------------- commentary grading

def build_mini(tmp):
    """The 6-month mini company of tests/fixtures/mini.py (no adjustments), as a company folder the CLI can run."""
    from tests.fixtures import mini
    n = 6
    mini.build(tmp, mini.months(n), [100000] * n, [40000] * n, [100] * n, [50] * n)
    return tmp / "companies" / "mini"


GOOD_TARGETS = {"good_abc_2026-06.md": "abc", "good_mini_no_adjustments.md": "mini"}


def grade_commentary(on_row=lambda row: None):
    rows = []
    with tempfile.TemporaryDirectory(prefix="evals_c_") as td:
        tmp = Path(td)
        runs = {}
        for key, cdir in (("abc", ABC_DIR), ("mini", None)):
            if cdir is None:
                cdir = build_mini(tmp / "mini_src")
            out = tmp / f"out_{key}"
            r = cli("run", "--company-dir", cdir, "--out", out)
            folders = [d for d in out.rglob("digest.json")]
            runs[key] = folders[0].parent if r.returncode == 0 and len(folders) == 1 else None

        def render(draft, key):
            run_dir = runs[key]
            if run_dir is None:
                return None, None, False
            (run_dir / "commentary.md").unlink(missing_ok=True)
            r = cli("render", "--run-dir", run_dir, "--draft", draft)
            rep = run_dir / "grounding_report.json"
            codes = sorted({f["code"] for f in json.loads(rep.read_text(encoding="utf-8"))["failures"]}) if rep.is_file() else None
            return r.returncode, codes, (run_dir / "commentary.md").is_file()

        for p in sorted(GOOD_DIR.glob("*.md")):
            key = GOOD_TARGETS.get(p.name)
            chk = Checks()
            if chk.add("known good draft", key is not None, f"{p.name} is not in GOOD_TARGETS in run_evals.py: say which run it belongs to"):
                code, codes, written = render(p, key)
                chk.add(f"{key} run available", code is not None, f"the {key} run failed, cannot render")
                if code is not None:
                    chk.eq("exit 0", code, 0)
                    chk.eq("no failure codes", codes, [])
                    chk.eq("commentary.md written", written, True)
            rows.append({"id": f"C-{p.stem}", "kind": "good draft", "expected": "pass", "checks": chk, "source": p})
            on_row(rows[-1])
        for p in sorted(BAD.glob("*.md")):
            m = re.match(r"g(\d+)_", p.name)
            chk = Checks()
            want = f"G{m.group(1)}" if m else None
            if chk.add("file name says its code", want is not None, f"{p.name} must be named g<N>_<what>.md"):
                code, codes, written = render(p, "abc")
                chk.eq("exit 4", code, 4)
                chk.eq(f"failure codes exactly [{want}]", codes, [want])
                chk.eq("commentary.md not written", written, False)
            rows.append({"id": f"C-{p.stem}", "kind": "bad draft", "expected": f"fail {want}", "checks": chk, "source": p})
            on_row(rows[-1])
    return rows


# ---------------------------------------------------------------- regression baseline

def snapshot(rows):
    """{row id: {"source_sha256", "result", "checks": {check name: "PASS"|"FAIL"}}} for the baseline and its diff."""
    out = {}
    for r in rows:
        checks = {}
        for i in r["checks"].items:
            name, k = i["check"], 2
            while name in checks:                                  # keep a repeated check name distinct
                name, k = f"{i['check']} #{k}", k + 1
            checks[name] = "PASS" if i["ok"] else "FAIL"
        out[r["id"]] = {"source_sha256": sha(r["source"]) if r.get("source") else None,
                        "result": "PASS" if r["checks"].ok else "FAIL", "checks": checks}
    return out


def baseline_diff(base_rows, now, partial):
    """-> list of lines, empty when the results match the baseline. `partial`: only some cases were selected, so
    golden cases missing from this run are not reported as removed."""
    lines = []
    for rid in sorted(set(base_rows) | set(now)):
        b, n = base_rows.get(rid), now.get(rid)
        if b is None:
            lines.append(f"NEW CASE      {rid} ({n['result']}, {len(n['checks'])} checks)")
            continue
        if n is None:
            if not (partial and rid[0] in "GX"):
                lines.append(f"REMOVED CASE  {rid} (was {b['result']}, {len(b['checks'])} checks)")
            continue
        if b["source_sha256"] != n["source_sha256"]:
            lines.append(f"SOURCE EDITED {rid}: its case file / draft changed since the baseline")
        if b["result"] != n["result"]:
            lines.append(f"RESULT        {rid}: {b['result']} -> {n['result']}")
        for c in sorted(set(b["checks"]) | set(n["checks"])):
            if c not in b["checks"]:
                lines.append(f"NEW CHECK     {rid}: {c} ({n['checks'][c]})")
            elif c not in n["checks"]:
                lines.append(f"REMOVED CHECK {rid}: {c} (was {b['checks'][c]})")
            elif b["checks"][c] != n["checks"][c]:
                lines.append(f"CHANGED       {rid}: {c}: {b['checks'][c]} -> {n['checks'][c]}")
    return lines


# ---------------------------------------------------------------- report

def _row_line(r, w):
    items = r["checks"].items
    what = r.get("name") or f"{r['kind']}, expect {r['expected']}"
    out = [f"{r['id'].ljust(w)}  {'PASS' if r['checks'].ok else 'FAIL':6}  {sum(i['ok'] for i in items):>2}/{len(items):<3}  {what}"]
    out += [f"{' ' * w}    - {i['check']}: {i['detail']}" for i in items if not i["ok"]]
    return "\n".join(out)


def parse_args(argv):
    ap = argparse.ArgumentParser(prog="run_evals.py", description="Run the eval suite (see the module docstring).")
    ap.add_argument("ids", nargs="*", help="case ids to run (default: all); commentary grading always runs")
    ap.add_argument("--cases-dir", type=Path, default=CASES, help="folder of case files (default: evals/cases)")
    ap.add_argument("--results", type=Path, default=RESULTS, help="where to write the results JSON")
    ap.add_argument("--baseline", type=Path, default=BASELINE, help="baseline to compare with")
    ap.add_argument("--no-baseline", action="store_true", help="do not compare with the baseline")
    ap.add_argument("--update-baseline", action="store_true",
                    help="print the diff, then rewrite the baseline from this run (full run, every row passing)")
    a = ap.parse_args(argv)
    if a.update_baseline and (a.ids or a.no_baseline):
        ap.error("--update-baseline needs a full run: no case ids and no --no-baseline")
    return a


def main(argv):
    try:
        args = parse_args(argv)
    except SystemExit as e:
        return e.code
    t0 = time.perf_counter()
    wanted = set(args.ids)
    cases = [(p, json.loads(p.read_text(encoding="utf-8"))) for p in sorted(args.cases_dir.glob("*.json"))]
    if wanted:
        unknown = wanted - {c["id"] for _, c in cases}
        if unknown:
            print(f"error: unknown case id(s) {sorted(unknown)}", file=sys.stderr)
            return 2
        cases = [(p, c) for p, c in cases if c["id"] in wanted]
    guarded_before = {str(p.relative_to(ROOT)): sha(p) for p in GUARDED}
    if SKILL_DIR != (ROOT / ".claude" / "skills" / "monthly-trend-analysis").resolve():
        print(f"skill under test: {SKILL_DIR}")

    def progress(r):
        print(f"[{time.perf_counter() - t0:6.1f}s] {r['id']:<26} {'PASS' if r['checks'].ok else 'FAIL'}", flush=True)

    rows = []
    for p, c in cases:
        rows.append({"id": c["id"], "kind": "golden case", "name": c["name"], "checks": run_case(c), "source": p})
        progress(rows[-1])
    rows += grade_commentary(progress)

    guard = Checks()
    for p in GUARDED:
        rel = str(p.relative_to(ROOT))
        guard.eq(f"{rel} unchanged", sha(p), guarded_before[rel])
    rows.append({"id": "GUARD", "kind": "guard", "name": "protected files unchanged", "checks": guard})
    progress(rows[-1])

    w = max(len(r["id"]) for r in rows)
    print(f"\n{'case'.ljust(w)}  result  checks  what")
    print(f"{'-' * w}  ------  ------  ----")
    for r in rows:
        print(_row_line(r, w))
    passed = sum(r["checks"].ok for r in rows)
    print(f"\n{passed}/{len(rows)} passed "
          f"({sum(len(r['checks'].items) for r in rows)} checks, "
          f"{sum(not i['ok'] for r in rows for i in r['checks'].items)} failed)")

    now = snapshot(rows)
    diff = None
    if not args.no_baseline:
        if args.baseline.is_file():
            base = json.loads(args.baseline.read_text(encoding="utf-8"))
            diff = baseline_diff(base["rows"], now, partial=bool(wanted))
            print(f"\nbaseline {args.baseline.name}: " + (f"{len(diff)} difference(s)" if diff else "no differences")
                  + (" (partial run: only the selected cases and the commentary rows are compared)" if wanted else ""))
            for line in diff:
                print("  " + line)
        else:
            diff = [f"no baseline at {args.baseline}"]
            print(f"\nbaseline: none at {args.baseline}; create it with --update-baseline")

    elapsed = time.perf_counter() - t0
    args.results.parent.mkdir(parents=True, exist_ok=True)
    git = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "HEAD"], capture_output=True, text=True).stdout.strip()
    doc = {"schema_version": 1, "git_sha": git or "unavailable", "selected": sorted(wanted) or "all",
           "skill_dir": str(SKILL_DIR), "elapsed_seconds": round(elapsed, 1),
           "summary": {"rows": len(rows), "passed": passed, "failed": len(rows) - passed},
           "baseline_differences": diff,
           "rows": [{"id": r["id"], "kind": r["kind"], "result": "PASS" if r["checks"].ok else "FAIL",
                     "checks": r["checks"].items} for r in rows]}
    args.results.write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    print(f"wrote {args.results}")

    if args.update_baseline:
        if passed != len(rows):
            print("baseline NOT updated: fix the failing rows first", file=sys.stderr)
            return 1
        args.baseline.write_text(json.dumps({"schema_version": 1, "rows": now}, indent=1, ensure_ascii=False) + "\n",
                                 encoding="utf-8", newline="\n")
        print(f"baseline updated: {args.baseline} now holds {len(now)} rows, "
              f"{sum(len(r['checks']) for r in now.values())} checks")
        diff = None
    print(f"elapsed {elapsed:.1f}s")
    if passed != len(rows):
        return 1
    return 6 if diff else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
