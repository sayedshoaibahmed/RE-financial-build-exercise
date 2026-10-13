"""Tiny workbooks built in code (openpyxl) in the ABC layout.

Rows are mutable lists [col_A, col_B, values] so tests can tamper with a
leaf without touching its control row. Values are chronological
(oldest -> newest); write() reverses them for newest-first sheets.
Control rows are written as literal numbers (openpyxl cannot cache formulas).
"""
import copy
from datetime import datetime

import openpyxl

PERIODS = ["2025-01", "2025-02", "2025-03"]
MONTH = {"2025-01": "Jan 2025", "2025-02": "Feb 2025", "2025-03": "Mar 2025"}
HEADER_ROW = 5
FIRST_ROW = 7
SHEET = "Income Statement"

SOURCE = {
    "type": "excel_pl",
    "path": "unused.xlsx",
    "sheet": SHEET,
    "header_row": HEADER_ROW,
    "section_label_col": "A",
    "account_label_col": "B",
    "period_header_format": "%b %Y",
    "sections": [
        {"header": "Income", "total": "Total Income", "kind": "income"},
        {"header": "Cost of Goods Sold", "total": "Total Cost of Goods Sold", "kind": "cogs"},
        {"header": "Operating Expenses", "total": "Total Operating Expenses", "kind": "opex"},
    ],
    "groups": [{"header": "Payroll", "total": "Total Payroll"}],
    "derived_rows": {"gross_profit": "Gross Profit", "operating_income": "Operating Income",
                     "net_income": "Net Income"},
}

OTHER_SECTIONS = [
    {"header": "Other Income", "total": "Total Other Income", "kind": "other_income"},
    {"header": "Other Expense", "total": "Total Other Expense", "kind": "other_expense"},
]


def source(**over):
    s = copy.deepcopy(SOURCE)
    s.update(over)
    return s


def _sum(*lists):
    return [sum(v) for v in zip(*lists)]


def default_rows(with_other=False):
    sales, setup = [1000, 1100, 1200], [50, 60, 70]
    hosting = [300, 300, 350]
    rent, mkt = [100, 100, 100], [10, 20, 30]
    salaries, taxes = [400, 400, 400], [40, 40, 40]
    t_inc = _sum(sales, setup)
    gp = [a - b for a, b in zip(t_inc, hosting)]
    t_pay = _sum(salaries, taxes)
    t_opex = _sum(rent, t_pay, mkt)
    opinc = [a - b for a, b in zip(gp, t_opex)]
    rows = [
        ["Income", None, None],
        [None, "Sales", sales],
        [None, "Fees/Misc: Setup", setup],
        ["Total Income", None, t_inc],
        ["Cost of Goods Sold", None, None],
        [None, "Hosting", hosting],
        ["Total Cost of Goods Sold", None, hosting],
        [None, "Gross Profit", gp],
        ["Operating Expenses", None, None],
        [None, "Rent", rent],
        [None, "Payroll", None],
        [None, "Salaries", salaries],
        [None, "Payroll Taxes", taxes],
        [None, "Total Payroll", t_pay],
        [None, "Marketing", mkt],
        ["Total Operating Expenses", None, t_opex],
        [None, "Operating Income", opinc],
    ]
    net = opinc
    if with_other:
        oi, oe = [5, 5, 5], [2, 2, 2]
        rows += [
            ["Other Income", None, None], [None, "Interest", oi], ["Total Other Income", None, oi],
            ["Other Expense", None, None], [None, "Fees", oe], ["Total Other Expense", None, oe],
        ]
        net = [a + b - c for a, b, c in zip(opinc, oi, oe)]
    rows.append([None, "Net Income", net])
    return rows


def row_of(rows, label):
    """1-based sheet row of a label (as written by write())."""
    for i, (a, b, _) in enumerate(rows):
        if label in (a, b):
            return FIRST_ROW + i
    raise KeyError(label)


def write(path, rows, order="oldest", periods=PERIODS, header_values=None, extra_empty_cols=3):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.cell(1, 1, "Title that is misleading: month ended Dec 31, 2099")
    ws.cell(HEADER_ROW, 2, "Account")
    ordered = list(periods) if order == "oldest" else list(reversed(periods))
    heads = header_values or [datetime.strptime(p, "%Y-%m").strftime("%b %Y") for p in ordered]
    for j, h in enumerate(heads):
        ws.cell(HEADER_ROW, 3 + j, h)
    for k in range(extra_empty_cols):  # used range wider than the periods
        for rr in (HEADER_ROW, FIRST_ROW):  # styled-but-empty cells persist on save, like ABC's U..Y
            ws.cell(rr, 3 + len(heads) + k).number_format = "0.00"
    for i, (a, b, vals) in enumerate(rows):
        r = FIRST_ROW + i
        if a is not None:
            ws.cell(r, 1, a)
        if b is not None:
            ws.cell(r, 2, b)
        if vals is not None:
            vs = list(vals) if order == "oldest" else list(reversed(vals))
            for j, v in enumerate(vs):
                if v is not None:
                    ws.cell(r, 3 + j, v)
    wb.save(path)
    return path
