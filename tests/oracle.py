"""Independent, deliberately naive scan of the ABC workbook.

Shares NO code with the trend package. It trusts nothing but openpyxl and
the visible layout: header row 5, labels in columns A/B, numbers in the
period columns. Slow and dumb on purpose; tests use it to derive structural
facts (row counts, label sets, periods) instead of hard-coding them.
"""
from datetime import datetime
from pathlib import Path

import openpyxl

WORKBOOK = Path(__file__).resolve().parent.parent / "data" / "ABC Inc - Financial Statements.xlsx"
SHEET = "Income Statement"
DERIVED = {"Gross Profit", "Operating Income", "Net Income"}
SECTIONS = {"Income": "income", "Cost of Goods Sold": "cogs", "Operating Expenses": "opex"}


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def scan(path=WORKBOOK):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET]

    # periods: every header-row cell that parses as "Mon YYYY"
    periods = []
    period_cols = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(5, c).value
        if isinstance(v, str):
            try:
                d = datetime.strptime(v.strip(), "%b %Y")
            except ValueError:
                continue
            periods.append(d.strftime("%Y-%m"))
            period_cols.append(c)

    leaves, subtotals, derived, headers = [], [], [], []
    cur = None
    for r in range(6, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        label = a if a else b
        if label is None:
            continue
        label = str(label).strip()
        vals = [ws.cell(r, c).value for c in period_cols]
        all_num = all(_is_num(v) for v in vals)
        any_num = any(_is_num(v) for v in vals)
        if not any_num:
            headers.append((r, label))
            if label in SECTIONS:
                cur = SECTIONS[label]
        elif label in DERIVED:
            derived.append((r, label))
        elif label.startswith("Total"):
            subtotals.append((r, label))
        elif all_num:
            leaves.append({"row": r, "label": label, "values": vals, "section": cur})
        else:
            raise AssertionError(f"oracle: partial numeric row {r} {label!r}")

    return {
        "periods": periods,
        "period_cols": period_cols,
        "leaves": leaves,
        "leaf_labels": [x["label"] for x in leaves],
        "subtotals": subtotals,
        "derived": derived,
        "headers": headers,
        "slash_labels": [x["label"] for x in leaves if "/" in x["label"]],
        "all_zero_leaves": [x["label"] for x in leaves if all(v == 0 for v in x["values"])],
    }


# ---------------------------------------------------------------------------
# Naive metric calculations (P4). Independent of the trend package: plain loops,
# exact Fractions, no Decimal, no shared helpers. Percent values are in percent units.
# ---------------------------------------------------------------------------
from fractions import Fraction  # noqa: E402

import yaml  # noqa: E402


def cents(v):
    """Workbook float -> exact 2dp Fraction."""
    return Fraction(round(Fraction(repr(v)) * 100), 100)


def shift(period, k):
    y, m = int(period[:4]), int(period[5:7])
    n = y * 12 + (m - 1) + k
    return f"{n // 12:04d}-{n % 12 + 1:02d}"


def monthly(o, zeroed=()):
    """{period: {rev, cogs, gp, opex, opinc}} as Fractions. `zeroed` = {(account, period)}."""
    out = {}
    for i, p in enumerate(o["periods"]):
        t = {"income": Fraction(0), "cogs": Fraction(0), "opex": Fraction(0)}
        for leaf in o["leaves"]:
            if (leaf["label"], p) in zeroed:
                continue
            t[leaf["section"]] += cents(leaf["values"][i])
        out[p] = {"rev": t["income"], "cogs": t["cogs"], "gp": t["income"] - t["cogs"],
                  "opex": t["opex"], "opinc": t["income"] - t["cogs"] - t["opex"]}
    return out


def by_category(o, mapping_path, field, zeroed=()):
    """{period: {category: Fraction}} over opex leaves, using mapping.yaml read raw with PyYAML."""
    entries = {e["account"]: e for e in yaml.safe_load(Path(mapping_path).read_text(encoding="utf-8"))}
    out = {p: {} for p in o["periods"]}
    for leaf in o["leaves"]:
        if leaf["section"] != "opex":
            continue
        cat = entries[leaf["label"]][field]
        for i, p in enumerate(o["periods"]):
            if (leaf["label"], p) in zeroed:
                continue
            out[p][cat] = out[p].get(cat, Fraction(0)) + cents(leaf["values"][i])
    return out


def window(series, end, n):
    """Sum of series[...] over the n calendar months ending `end`; None if any month is absent."""
    months = [shift(end, -i) for i in range(n)]
    if any(m not in series for m in months):
        return None
    return sum((series[m] for m in months), Fraction(0))


def pct(num, den):
    return None if den <= 0 else num / den * 100
