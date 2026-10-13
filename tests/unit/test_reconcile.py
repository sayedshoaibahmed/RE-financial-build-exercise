import dataclasses
import re
import zipfile
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl.utils import get_column_letter

from tests import oracle
from tests.fixtures import builders as b
from trend.adapters.excel_pl import ExcelPLAdapter
from trend.config import load_company, parse_source
from trend.reconcile import BlockedError, ControlResult, ReconResult, ensure_ok, reconcile

ROOT = Path(__file__).resolve().parents[2]
D = Decimal


def tiny(tmp_path, rows=None, with_other=False, tamper=None):
    """Load a tiny workbook; `tamper(rows)` edits leaf values AFTER controls were computed."""
    rows = rows if rows is not None else b.default_rows(with_other=with_other)
    if tamper:
        tamper(rows)
    p = tmp_path / "t.xlsx"
    b.write(p, rows)
    src = b.source(path=str(p))
    if with_other:
        src["sections"] = src["sections"] + b.OTHER_SECTIONS
    cfg = parse_source(src, base_dir=tmp_path)
    records, controls = ExcelPLAdapter(cfg, "tiny").load()
    return cfg, records, controls


def fails(res):
    return {(r.control_name, r.period) for r in res.results if not r.ok}


def leaf(rows, label):
    return next(r for r in rows if r[1] == label)


# ---------- matching fixture passes ----------

def test_matching_fixture_passes_all_controls(tmp_path):
    cfg, records, controls = tiny(tmp_path)
    res = reconcile(records, controls, cfg)
    assert isinstance(res, ReconResult) and res.status == "OK"
    assert len(res.results) == 7 * len(b.PERIODS)
    assert all(isinstance(r, ControlResult) and r.ok for r in res.results)
    assert res.failures == []
    ensure_ok(res)  # does not raise


def test_result_fields(tmp_path):
    cfg, records, controls = tiny(tmp_path)
    res = reconcile(records, controls, cfg)
    r = next(x for x in res.results if x.control_id == "C1" and x.period == "2025-02")
    assert r.control_name == "Total Income"
    assert r.expected == D("1160.00") and r.actual == D("1160.00") and r.difference == D("0.00")
    assert r.source_ref == "Income Statement!D10"  # row of Total Income, column of Feb 2025
    ids = {x.control_id for x in res.results}
    assert ids == {"C1", "C2", "C3", "C4", "C6", "C7", "C8"}  # one group in this fixture


def test_two_groups_get_c4_and_c5(tmp_path):
    cfg, _, _ = tiny(tmp_path)
    src = b.source(path="x.xlsx")
    src["groups"] = src["groups"] + [{"header": "Mkt", "total": "Total Mkt"}]
    rows = b.default_rows()
    i = b.row_of(rows, "Marketing") - b.FIRST_ROW
    rows[i:i + 1] = [[None, "Mkt", None], [None, "Ads", [10, 20, 30]], [None, "Total Mkt", [10, 20, 30]]]
    p = tmp_path / "g.xlsx"
    b.write(p, rows)
    src["path"] = str(p)
    cfg2 = parse_source(src, base_dir=tmp_path)
    records, controls = ExcelPLAdapter(cfg2, "tiny").load()
    res = reconcile(records, controls, cfg2)
    assert res.status == "OK"
    assert {x.control_name: x.control_id for x in res.results if x.control_id in ("C4", "C5")} == \
        {"Total Payroll": "C4", "Total Mkt": "C5"}


# ---------- tampered leaf ----------

def test_tampered_income_leaf_fails_c1_c3_c7_c8_with_exact_difference(tmp_path):
    def t(rows):
        leaf(rows, "Sales")[2] = [1000, 1100 + 1000, 1200]  # Feb +1000
    cfg, records, controls = tiny(tmp_path, tamper=t)
    res = reconcile(records, controls, cfg)
    assert res.status == "BLOCKED"
    assert fails(res) == {("Total Income", "2025-02"), ("Gross Profit", "2025-02"),
                          ("Operating Income", "2025-02"), ("Net Income", "2025-02")}
    for r in res.failures:
        assert r.difference == D("1000.00")  # actual (recomputed) - expected (workbook)
        assert r.actual - r.expected == r.difference
    ids = {r.control_id for r in res.failures}
    assert ids == {"C1", "C3", "C7", "C8"}
    gp = next(r for r in res.failures if r.control_id == "C3")
    assert gp.source_ref == f"Income Statement!D{b.row_of(b.default_rows(), 'Gross Profit')}"


def test_tampered_opex_group_leaf_fails_c4_c6_c7_c8(tmp_path):
    def t(rows):
        leaf(rows, "Salaries")[2] = [400, 400, 400 - 25]  # Mar -25
    cfg, records, controls = tiny(tmp_path, tamper=t)
    res = reconcile(records, controls, cfg)
    assert {r.control_id for r in res.failures} == {"C4", "C6", "C7", "C8"}
    assert {r.period for r in res.failures} == {"2025-03"}
    by = {r.control_id: r.difference for r in res.failures}
    assert by["C4"] == D("-25.00") and by["C6"] == D("-25.00")
    assert by["C7"] == D("25.00") and by["C8"] == D("25.00")  # opex lower -> op income higher


def test_tampered_cogs_leaf_fails_c2_c3_c7_c8(tmp_path):
    def t(rows):
        leaf(rows, "Hosting")[2] = [301, 300, 350]
    cfg, records, controls = tiny(tmp_path, tamper=t)
    res = reconcile(records, controls, cfg)
    assert {r.control_id for r in res.failures} == {"C2", "C3", "C7", "C8"}


# ---------- tolerance edges ----------

def _with_income_control(controls, period, delta):
    out = []
    for c in controls:
        if c.control_name == "Total Income" and c.period == period:
            c = dataclasses.replace(c, amount=c.amount + D(delta))
        out.append(c)
    return out


@pytest.mark.parametrize("delta,ok", [
    ("0.009", True), ("-0.009", True), ("0.010", True), ("-0.010", True),
    ("0.011", False), ("-0.011", False), ("0", True),
])
def test_default_tolerance_is_one_cent(tmp_path, delta, ok):
    cfg, records, controls = tiny(tmp_path)
    res = reconcile(records, _with_income_control(controls, "2025-02", delta), cfg)
    c1 = next(r for r in res.results if r.control_id == "C1" and r.period == "2025-02")
    assert c1.ok is ok
    assert (res.status == "OK") is ok
    assert c1.difference == -D(delta)  # actual - expected


def test_tolerance_is_configurable(tmp_path):
    cfg, records, controls = tiny(tmp_path)
    c = _with_income_control(controls, "2025-02", "0.011")
    assert reconcile(records, c, cfg, tolerance=D("0.02")).status == "OK"
    c = _with_income_control(controls, "2025-02", "0.009")
    assert reconcile(records, c, cfg, tolerance=D("0")).status == "BLOCKED"


def test_negative_tolerance_rejected(tmp_path):
    cfg, records, controls = tiny(tmp_path)
    with pytest.raises(ValueError):
        reconcile(records, controls, cfg, tolerance=D("-0.01"))


# ---------- other sections (C8) and missing pieces ----------

def test_c8_includes_other_sections(tmp_path):
    cfg, records, controls = tiny(tmp_path, with_other=True)
    res = reconcile(records, controls, cfg)
    assert res.status == "OK"
    ni = next(r for r in res.results if r.control_id == "C8" and r.period == "2025-01")
    opinc = next(r for r in res.results if r.control_id == "C7" and r.period == "2025-01")
    assert ni.actual == opinc.actual + D("5.00") - D("2.00")


def test_c8_detects_tampered_other_income(tmp_path):
    def t(rows):
        leaf(rows, "Interest")[2] = [5, 5, 9]
    cfg, records, controls = tiny(tmp_path, with_other=True, tamper=t)
    res = reconcile(records, controls, cfg)
    assert fails(res) == {("Total Other Income", "2025-03"), ("Net Income", "2025-03")}


def test_missing_control_cell_fails_closed(tmp_path):
    cfg, records, controls = tiny(tmp_path)
    kept = [c for c in controls if not (c.control_name == "Gross Profit" and c.period == "2025-01")]
    res = reconcile(records, kept, cfg)
    assert res.status == "BLOCKED"
    r = next(x for x in res.failures)
    assert (r.control_name, r.period) == ("Gross Profit", "2025-01")
    assert r.expected is None and r.difference is None and r.reason == "MISSING_CONTROL"


def test_blocked_raises_with_recon_fail_and_blocks_downstream(tmp_path):
    def t(rows):
        leaf(rows, "Sales")[2] = [1, 1100, 1200]
    cfg, records, controls = tiny(tmp_path, tamper=t)
    res = reconcile(records, controls, cfg)
    with pytest.raises(BlockedError) as e:
        ensure_ok(res)
    assert e.value.reason == "RECON_FAIL"
    assert e.value.result is res
    assert "Income Statement!C" in str(e.value)  # cell reference is in the message


def test_reconcile_does_not_mutate_inputs(tmp_path):
    cfg, records, controls = tiny(tmp_path)
    r0, c0 = list(records), list(controls)
    reconcile(records, controls, cfg)
    assert records == r0 and controls == c0


# ---------- real ABC workbook ----------

@pytest.fixture(scope="module")
def abc():
    company = load_company(ROOT / "companies" / "abc_inc")
    records, controls = ExcelPLAdapter(company.source, company.company_id).load()
    return company, records, controls, oracle.scan()


def test_abc_all_controls_pass_in_all_periods(abc):
    company, records, controls, o = abc
    res = reconcile(records, controls, company.source)
    assert res.status == "OK", [(r.control_id, r.period, r.difference) for r in res.failures]
    n_controls = len(o["subtotals"]) + len(o["derived"])
    assert len(res.results) == n_controls * len(o["periods"])
    assert {r.period for r in res.results} == set(o["periods"])
    assert {r.control_id for r in res.results} == {f"C{i}" for i in range(1, 9)}
    assert all(r.ok for r in res.results)


# ---------- tampered copy of the real workbook (temp folder only) ----------

def tamper_xlsx(src: Path, dst: Path, cell: str, delta: str):
    """Edit one literal numeric cell in sheet1.xml, keeping every cached formula value."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                pat = re.compile(r'(<c r="%s"(?: [^>]*)?><v>)([^<]*)(</v>)' % cell)
                assert pat.search(xml), f"{cell} is not a literal numeric cell"
                xml = pat.sub(lambda m: m.group(1) + str(D(m.group(2)) + D(delta)) + m.group(3), xml, count=1)
                data = xml.encode("utf-8")
            zout.writestr(item, data)


def _cell_of(o, label, period):
    leaf_ = next(x for x in o["leaves"] if x["label"] == label)
    col = o["period_cols"][o["periods"].index(period)]
    return f"{get_column_letter(col)}{leaf_['row']}", col


def _control_ref(o, label, col):
    row = next(r for r, l in o["subtotals"] + o["derived"] if l == label)
    return f"Income Statement!{get_column_letter(col)}{row}"


def test_tampered_abc_copy_is_blocked_with_cell_references(abc, tmp_path):
    company, _, _, o = abc
    period = o["periods"][0]
    cell, col = _cell_of(o, "Usage Fees", period)
    copy = tmp_path / "tampered.xlsx"
    tamper_xlsx(company.source.path, copy, cell, "1000.00")
    src = dataclasses.replace(company.source, path=copy)
    records, controls = ExcelPLAdapter(src, "abc_inc").load()
    res = reconcile(records, controls, src)
    assert res.status == "BLOCKED"
    assert {(r.control_name, r.period) for r in res.failures} == {
        (n, period) for n in ("Total Income", "Gross Profit", "Operating Income", "Net Income")}
    assert {r.difference for r in res.failures} == {D("1000.00")}
    assert {r.source_ref for r in res.failures} == {
        _control_ref(o, n, col) for n in ("Total Income", "Gross Profit", "Operating Income", "Net Income")}
    with pytest.raises(BlockedError) as e:
        ensure_ok(res)
    assert _control_ref(o, "Total Income", col) in str(e.value)


def test_tampered_abc_group_leaf_hits_group_and_opex_controls(abc, tmp_path):
    company, _, _, o = abc
    period = o["periods"][3]
    cell, col = _cell_of(o, "Events & Conferences", period)
    copy = tmp_path / "tampered2.xlsx"
    tamper_xlsx(company.source.path, copy, cell, "-250.50")
    src = dataclasses.replace(company.source, path=copy)
    records, controls = ExcelPLAdapter(src, "abc_inc").load()
    res = reconcile(records, controls, src)
    assert {r.period for r in res.failures} == {period}
    assert {r.control_name for r in res.failures} == {
        "Total Marketing", "Total Operating Expenses", "Operating Income", "Net Income"}
    assert {r.control_id for r in res.failures} == {"C5", "C6", "C7", "C8"}


def test_original_workbook_is_untouched_by_the_tests(abc):
    company, *_ = abc
    assert company.source.path == oracle.WORKBOOK.resolve() or company.source.path.samefile(oracle.WORKBOOK)
