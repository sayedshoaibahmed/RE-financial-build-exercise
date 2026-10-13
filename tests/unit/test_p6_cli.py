"""Stage C: the run / check commands, exit codes, BLOCKED handling, determinism and an end-to-end ABC run via the CLI."""
import hashlib
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

from tests.fixtures import mini
from tests.fixtures.mini import months
from trend.digest import SIZE_BUDGET_BYTES
from trend.pipeline import EXIT_BLOCKED, EXIT_OK, EXIT_USAGE, default_out_base, run as run_pipeline

ROOT = Path(__file__).resolve().parents[2]
SKILL = ROOT / ".claude" / "skills" / "monthly-trend-analysis"
CLI = SKILL / "scripts" / "trend_cli.py"
ABC_DIR = ROOT / "companies" / "abc_inc"
RUN_ID = "2026-06_w12-2026-06"


def cli(*args, cwd=ROOT):
    return subprocess.run([sys.executable, str(CLI), *map(str, args)], cwd=cwd, capture_output=True, text=True)


def tree(*dirs):
    out = {}
    for d in dirs:
        for p in sorted(Path(d).rglob("*")):
            if p.is_file() and "__pycache__" not in p.parts:
                out[str(p)] = hashlib.sha256(p.read_bytes()).hexdigest()
    return out


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ---------------------------------------------------------------- end to end on the real ABC workbook, twice

@pytest.fixture(scope="session")
def abc_runs(tmp_path_factory):
    before = tree(ABC_DIR, SKILL / "reference", SKILL / "templates") if (SKILL / "templates").exists() else tree(ABC_DIR, SKILL / "reference")
    roots = [tmp_path_factory.mktemp("run1"), tmp_path_factory.mktemp("run2")]
    procs = [cli("run", "--company", "abc_inc", "--out", r) for r in roots]
    after = tree(ABC_DIR, SKILL / "reference", SKILL / "templates") if (SKILL / "templates").exists() else tree(ABC_DIR, SKILL / "reference")
    return procs, [r / "abc_inc" / RUN_ID for r in roots], before, after


def test_cli_end_to_end_on_the_real_abc_workbook(abc_runs):
    (p1, p2), (d1, d2), *_ = abc_runs
    assert p1.returncode == EXIT_OK == p2.returncode, p1.stderr
    assert sorted(f.name for f in d1.iterdir()) == ["digest.json", "metrics.json", "run_manifest.json", "trend_pack.xlsx"]
    for name in ("digest.json", "metrics.json", "trend_pack.xlsx"):
        assert name in p1.stdout and (d1 / name).stat().st_size > 1000
    assert "OK" in p1.stdout


def test_end_to_end_outputs_have_the_anchors(abc_runs):
    _, (d1, _), *_ = abc_runs
    digest = json.loads((d1 / "digest.json").read_text(encoding="utf-8"))
    metrics = json.loads((d1 / "metrics.json").read_text(encoding="utf-8"))
    f = digest["allowed_facts"]
    head = {r["label"]: r for r in digest["headline"]["rows"]}
    assert f[head["Revenue"]["rep"]["value"]] == "195,898.16"
    assert metrics["facts"]["rep.rev.total.2026-06"] == "195898.16"
    assert metrics["facts"]["rep.opinc.2025-12"] == "534551.41" and metrics["facts"]["adj.opinc.2025-12"] == "-375217.90"
    assert metrics["facts"]["rep.rev.total.tw12.2026-06"] == "2303975.51"
    assert metrics["status"] == "OK" and metrics["reconciliation"]["status"] == "OK"
    assert len(metrics["reconciliation"]["results"]) == 180
    assert metrics["meta"]["input_sha256"] == sha(ROOT / "data" / "ABC Inc - Financial Statements.xlsx")
    assert digest["input_sha256"] == metrics["meta"]["input_sha256"] and digest["config_sha256"] == metrics["meta"]["config_sha256"]
    assert metrics["flags"]["counts"]["NEGATIVE_VALUE"] == 9 and metrics["flags"]["counts"]["NONRECURRING"] == 2


def test_digest_on_disk_is_within_the_size_budget(abc_runs):
    _, (d1, _), *_ = abc_runs
    size = (d1 / "digest.json").stat().st_size
    assert size <= SIZE_BUDGET_BYTES, f"{size} > {SIZE_BUDGET_BYTES}"
    assert size < (d1 / "metrics.json").stat().st_size / 10


def test_workbook_data_sheet_ids_all_exist_in_metrics_json(abc_runs):
    _, (d1, _), *_ = abc_runs
    metrics = json.loads((d1 / "metrics.json").read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(d1 / "trend_pack.xlsx", read_only=True)
    ids = [r[0] for r in wb["Data"].iter_rows(min_row=2, values_only=True)]
    wb.close()
    assert len(ids) == len(metrics["facts"]) and set(ids) == set(metrics["facts"])


def test_two_cli_runs_give_byte_identical_json_and_identical_cells(abc_runs):
    _, (d1, d2), *_ = abc_runs
    for name in ("digest.json", "metrics.json"):
        assert (d1 / name).read_bytes() == (d2 / name).read_bytes(), name
    w1 = openpyxl.load_workbook(d1 / "trend_pack.xlsx", read_only=True)
    w2 = openpyxl.load_workbook(d2 / "trend_pack.xlsx", read_only=True)
    assert w1.sheetnames == w2.sheetnames
    for a, b in zip(w1.worksheets, w2.worksheets):
        assert list(a.iter_rows(values_only=True)) == list(b.iter_rows(values_only=True)), a.title
    w1.close(), w2.close()


def test_the_run_wrote_nothing_into_companies_or_the_skill_folder(abc_runs):
    *_, before, after = abc_runs
    assert before == after


def test_blocked_flag_free_success_run_has_no_blocked_report(abc_runs):
    _, (d1, _), *_ = abc_runs
    assert not (d1 / "blocked_report.json").exists() and not (d1 / "unmapped_accounts.yaml").exists()


def test_cli_window_options_change_the_run_id_and_the_windows(tmp_path):
    r = cli("run", "--company", "abc_inc", "--window", "6", "--window-end", "2026-03", "--out", tmp_path)
    assert r.returncode == 0, r.stderr
    d = tmp_path / "abc_inc" / "2026-03_w6-2026-03"                      # as-of follows the window end
    digest = json.loads((d / "digest.json").read_text(encoding="utf-8"))
    assert digest["as_of"] == "2026-03"
    (w,) = digest["windows"]
    f = digest["allowed_facts"]
    rows = {x["label"]: x for x in w["rows"]}
    assert (w["length"], w["end"]) == (6, "2026-03")
    assert f[rows["Revenue"]["rep"]["value"]] == "1,148,039.20"
    assert f[rows["Operating income"]["rep"]["value"]] == "-1,095,637.76" and f[rows["Operating income"]["adj"]["value"]] == "-2,005,407.07"


def test_cli_period_sugar_and_as_of(tmp_path):
    r = cli("run", "--company-dir", ABC_DIR, "--period", "2025-10:2026-03", "--as-of", "2026-04", "--out", tmp_path)
    assert r.returncode == 0, r.stderr
    d = tmp_path / "abc_inc" / "2026-04_w6-2026-03"
    digest = json.loads((d / "digest.json").read_text(encoding="utf-8"))
    assert digest["as_of"] == "2026-04" and digest["window_lengths"] == [6] and digest["window_end"] == "2026-03"
    metrics = json.loads((d / "metrics.json").read_text(encoding="utf-8"))
    assert not [k for k in metrics["facts"] if k.endswith("2026-05") or k.endswith("2026-06")]     # no lookahead


def _digest(d):
    return json.loads((d / "digest.json").read_text(encoding="utf-8"))


def test_window_end_alone_sets_the_as_of_month(tmp_path):
    r = cli("run", "--company", "abc_inc", "--window", "6", "--window-end", "2026-03", "--out", tmp_path)
    assert r.returncode == 0, r.stderr
    d = tmp_path / "abc_inc" / "2026-03_w6-2026-03"
    dg = _digest(d)
    assert dg["as_of"] == "2026-03" and dg["window_end"] == "2026-03"
    metrics = json.loads((d / "metrics.json").read_text(encoding="utf-8"))
    assert not [k for k in metrics["facts"] if k.endswith("2026-04") or k.endswith("2026-05") or k.endswith("2026-06")]
    f = dg["allowed_facts"]
    (w,) = dg["windows"]
    rows = {x["label"]: x for x in w["rows"]}
    assert f[rows["Revenue"]["rep"]["value"]] == "1,148,039.20"
    assert f[rows["Operating income"]["rep"]["value"]] == "-1,095,637.76"
    assert f[rows["Operating income"]["adj"]["value"]] == "-2,005,407.07"


def test_period_sets_the_as_of_month_to_its_end(tmp_path):
    r = cli("run", "--company", "abc_inc", "--period", "2025-10:2026-03", "--out", tmp_path)
    assert r.returncode == 0, r.stderr
    assert _digest(tmp_path / "abc_inc" / "2026-03_w6-2026-03")["as_of"] == "2026-03"


def test_explicit_as_of_wins_over_the_window_end(tmp_path):
    r = cli("run", "--company", "abc_inc", "--window", "6", "--window-end", "2026-03", "--as-of", "2026-06", "--out", tmp_path)
    assert r.returncode == 0, r.stderr
    dg = _digest(tmp_path / "abc_inc" / "2026-06_w6-2026-03")
    assert dg["as_of"] == "2026-06" and dg["window_end"] == "2026-03"


def test_several_window_lengths_with_one_window_end_share_one_as_of(tmp_path):
    r = cli("run", "--company", "abc_inc", "--window", "3", "--window", "6", "--window-end", "2026-03", "--out", tmp_path)
    assert r.returncode == 0, r.stderr
    assert _digest(tmp_path / "abc_inc" / "2026-03_w3+6-2026-03")["as_of"] == "2026-03"


def test_window_end_later_than_the_data_still_errors(tmp_path):
    r = cli("run", "--company", "abc_inc", "--window-end", "2030-01", "--out", tmp_path)
    assert r.returncode == EXIT_USAGE and "WINDOW_END_AFTER_DATA" in r.stderr
    assert not (tmp_path / "abc_inc").exists()


def test_run_help_documents_the_as_of_default():
    flat = " ".join(cli("run", "--help").stdout.split())
    assert "also the as-of month unless --as-of is given" in flat
    assert "default: the --window-end / --period end if given" in flat


# ---------------------------------------------------------------- the run folder name and library-level runs

def good(tmp_path, n=6, **kw):
    mini.build(tmp_path, months(n), [100000] * n, [40000] * n, [100] * n, [50] * n, **kw)
    return tmp_path / "companies" / "mini"


def test_run_id_and_default_output_base(tmp_path):
    cdir = good(tmp_path)
    res = run_pipeline(cdir, out_base=tmp_path / "o")
    assert res.status == "OK" and res.exit_code == EXIT_OK
    assert res.out_dir == tmp_path / "o" / "mini" / "2025-06_w12-2025-06"
    assert sorted(p.name for p in res.files) == ["digest.json", "metrics.json", "run_manifest.json", "trend_pack.xlsx"]
    assert default_out_base(cdir) == tmp_path / "outputs"            # PLAN: outputs/<company_id>/<run_id>/
    assert default_out_base(ABC_DIR) == ROOT / "outputs"
    res2 = run_pipeline(cdir, out_base=tmp_path / "o", window_lengths=(3, 6), window_end="2025-05")
    assert res2.out_dir.name == "2025-05_w3+6-2025-05"
    res3 = run_pipeline(cdir, out_base=tmp_path / "o", window_lengths=(3, 6), window_end="2025-05", as_of="2025-06")
    assert res3.out_dir.name == "2025-06_w3+6-2025-05"                      # an explicit as-of wins


# ---------------------------------------------------------------- BLOCKED: only the blocked report is written

def blocked(tmp_path):
    n = 6
    mini.build(tmp_path, months(n), [100000] * n, [40000] * n, [100] * n, [50] * n,
               sales_leaf=[100000, 100000, 100001, 100000, 100000, 100000])      # fails C1 (and C3, C7, C8) in 2025-03
    return tmp_path / "companies" / "mini"


def test_blocked_run_writes_only_the_blocked_report(tmp_path):
    cdir = blocked(tmp_path)
    res = run_pipeline(cdir, out_base=tmp_path / "o")
    assert res.status == "BLOCKED" and res.exit_code == EXIT_BLOCKED
    run_dir = tmp_path / "o" / "mini" / "latest_w12-latest"
    assert sorted(p.name for p in run_dir.iterdir()) == ["blocked_report.json"]
    assert [p.name for p in res.files] == ["blocked_report.json"]
    rep = json.loads((run_dir / "blocked_report.json").read_text(encoding="utf-8"))
    assert rep["status"] == "BLOCKED" and rep["reason"] == "RECON_FAIL" and rep["company_id"] == "mini"
    c1 = next(c for c in rep["failed_controls"] if c["control_id"] == "C1")
    assert (c1["period"], c1["expected"], c1["actual"], c1["difference"]) == ("2025-03", "100000.00", "100001.00", "1.00")
    assert re.fullmatch(r"Income Statement!E\d+", c1["source_ref"])
    assert {c["control_id"] for c in rep["failed_controls"]} == {"C1", "C3", "C7", "C8"}
    assert rep["flag"]["type"] == "BLOCKED" and rep["flag"]["observed"] == "RECON_FAIL"
    for forbidden in ("metrics", "flags_summary", "headline", "facts", "digest"):
        assert forbidden not in rep


def test_blocked_run_removes_stale_outputs_of_the_same_run_folder(tmp_path):
    cdir = blocked(tmp_path)
    run_dir = tmp_path / "o" / "mini" / "latest_w12-latest"
    run_dir.mkdir(parents=True)
    for n in ("metrics.json", "digest.json", "trend_pack.xlsx", "run_manifest.json"):
        (run_dir / n).write_text("stale")
    run_pipeline(cdir, out_base=tmp_path / "o")
    assert sorted(p.name for p in run_dir.iterdir()) == ["blocked_report.json"]


def test_unmapped_account_blocks_and_writes_the_unmapped_list_beside_the_report(tmp_path):
    cdir = good(tmp_path)
    ents = [e for e in yaml.safe_load((cdir / "mapping.yaml").read_text()) if e["account"] != "Rent"]
    (cdir / "mapping.yaml").write_text(yaml.safe_dump(ents))
    res = run_pipeline(cdir, out_base=tmp_path / "o")
    assert res.status == "BLOCKED" and res.exit_code == EXIT_BLOCKED
    names = sorted(p.name for p in (tmp_path / "o" / "mini" / "latest_w12-latest").iterdir())
    assert names == ["blocked_report.json", "unmapped_accounts.yaml"]
    rep = json.loads((tmp_path / "o" / "mini" / "latest_w12-latest" / "blocked_report.json").read_text())
    assert rep["reason"] == "UNMAPPED_ACCOUNT" and rep["failed_controls"] == []


def test_adapter_error_blocks(tmp_path):
    cdir = good(tmp_path)
    import openpyxl as ox
    xl = tmp_path / "mini.xlsx"
    wb = ox.load_workbook(xl)
    ws = wb.active
    ws.cell(9, 2, "Total Widgets")                                         # an undeclared total in a data row
    wb.save(xl)
    res = run_pipeline(cdir, out_base=tmp_path / "o")
    assert res.status == "BLOCKED"
    rep = json.loads(next(iter(res.files)).read_text())
    assert rep["reason"] in ("UNDECLARED_TOTAL", "NO_CACHED_VALUES", "UNCLASSIFIED_ROW", "NON_NUMERIC") and rep["source_ref"]


# ---------------------------------------------------------------- exit codes through the real CLI

def test_exit_code_ok_blocked_and_usage(tmp_path):
    ok = cli("run", "--company-dir", good(tmp_path / "ok"), "--out", tmp_path / "o1")
    assert ok.returncode == 0, ok.stderr
    bad = cli("run", "--company-dir", blocked(tmp_path / "bl"), "--out", tmp_path / "o2")
    assert bad.returncode == 3 and "BLOCKED" in bad.stderr and "RECON_FAIL" in bad.stderr
    assert sorted(p.name for p in (tmp_path / "o2" / "mini" / "latest_w12-latest").iterdir()) == ["blocked_report.json"]
    assert (EXIT_OK, EXIT_BLOCKED, EXIT_USAGE) == (0, 3, 2)


@pytest.mark.parametrize("args,needle", [
    (["run"], "required"),
    (["run", "--company", "no_such_company"], "not found"),
    (["run", "--company", "../etc"], "invalid company"),
    (["run", "--company-dir", "x", "--company", "abc_inc"], "not allowed"),
    (["run", "--company", "abc_inc", "--window", "0"], "window"),
    (["run", "--company", "abc_inc", "--window", "x"], "invalid int"),
    (["run", "--company", "abc_inc", "--window-end", "soon"], "window_end"),
    (["run", "--company", "abc_inc", "--window-end", "2030-01"], "WINDOW_END_AFTER_DATA"),
    (["run", "--company", "abc_inc", "--as-of", "April"], "as_of"),
    (["run", "--company", "abc_inc", "--as-of", "2020-01"], "NO_DATA"),
    (["run", "--company", "abc_inc", "--period", "2026-06:2025-07"], "START"),
    (["run", "--company", "abc_inc", "--period", "2025-10:2026-03", "--window", "6"], "--period"),
    (["run", "--company", "abc_inc", "--window", "3", "--window", "3"], "window"),
    (["nonsense"], "invalid choice"),
    (["check"], "required"),
])
def test_usage_and_config_errors_exit_2_and_write_nothing(tmp_path, args, needle):
    r = cli(*args, "--out", tmp_path / "o") if args[0] == "run" else cli(*args)
    assert r.returncode == 2, (args, r.returncode, r.stderr)
    assert needle.lower() in r.stderr.lower(), r.stderr
    assert not (tmp_path / "o").exists()


def test_out_dir_inside_the_companies_folder_is_refused(tmp_path):
    cdir = good(tmp_path)                                  # a throwaway company: a broken guard cannot pollute the repo
    for bad in (cdir / "out", cdir.parent / "out", cdir.parent):
        r = cli("run", "--company-dir", cdir, "--out", bad)
        assert r.returncode == 2 and "outside" in r.stderr.lower(), bad
    assert sorted(p.name for p in cdir.parent.iterdir()) == ["mini"]       # nothing was created under companies/
    assert sorted(p.name for p in cdir.iterdir()) == ["company.yaml", "mapping.yaml"]


def test_out_dir_inside_the_skill_folder_is_refused(tmp_path):
    from trend.pipeline import UsageError, _check_out      # direct call: never runs a pipeline that could write there
    cdir = good(tmp_path)
    for bad in (SKILL / "out", SKILL / "scripts" / "out", SKILL):
        with pytest.raises(UsageError):
            _check_out(bad, cdir)
    assert _check_out(tmp_path / "ok", cdir) == (tmp_path / "ok").resolve()
    assert not (SKILL / "out").exists()


def test_exit_codes_are_documented_in_help():
    out = cli("run", "--help").stdout + cli("--help").stdout
    assert "Exit codes" in out
    for line in ("0  ", "2  ", "3  ", "4  "):
        assert line in out
    assert "BLOCKED" in out and "usage" in out.lower() and "config" in out.lower()
    assert "Exit codes" in CLI.read_text(encoding="utf-8")


# ---------------------------------------------------------------- check

def test_check_prints_the_reconciliation_table_and_writes_nothing():
    r = cli("check", "--company", "abc_inc")
    assert r.returncode == 0, r.stderr
    lines = [ln for ln in r.stdout.splitlines() if " | " in ln]
    assert lines[0].startswith("control | name | period | expected | actual | difference | status | source cell")
    assert len(lines) == 1 + 180
    assert any(ln.startswith("C9a | ") for ln in lines) and all(" | OK | " in ln for ln in lines[1:])
    assert "reconciliation: OK" in r.stdout and "180 controls" in r.stdout


def test_check_on_a_failing_workbook_exits_3_with_the_cell_reference(tmp_path):
    r = cli("check", "--company-dir", blocked(tmp_path))
    assert r.returncode == 3
    fails = [ln for ln in r.stdout.splitlines() if " | FAIL | " in ln]
    assert len(fails) == 4 and all(re.search(r"Income Statement!E\d+", ln) for ln in fails)
    assert "reconciliation: BLOCKED" in r.stdout
    assert not (tmp_path / "outputs").exists()


def test_check_reports_unmapped_accounts_as_blocked(tmp_path):
    cdir = good(tmp_path)
    (cdir / "mapping.yaml").unlink()
    r = cli("check", "--company-dir", cdir)
    assert r.returncode == 3 and "UNMAPPED_ACCOUNT" in r.stderr
    assert not (tmp_path / "outputs").exists()


# ---------------------------------------------------------------- the commands that are not part of P6

def test_render_without_a_draft_is_a_usage_error_not_a_stub(tmp_path):
    r = cli("render", "--company", "abc_inc", "--out", tmp_path)          # an empty output base: nothing to render
    assert r.returncode == 2 and "not implemented" not in r.stderr.lower() and "no run folder" in r.stderr.lower()


def test_list_accounts_still_works_with_company_id():
    r = cli("list-accounts", "--company", "abc_inc")
    assert r.returncode == 0 and "83 accounts, 83 mapped, 0 unmapped" in r.stdout
