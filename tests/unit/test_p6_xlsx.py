"""Stage B: the trend_pack.xlsx workbook, reopened with openpyxl and checked against the metrics and the oracle."""
import re
import zipfile
from decimal import Decimal
from fractions import Fraction

import openpyxl
import pytest

from tests import oracle
from trend.facts import display, month_name, parse_fact_id, unit_of
from trend.metrics import NA
from trend.writers.xlsx_pack import SHEETS, write_workbook

D = Decimal
META = {"input_sha256": "ab" * 32, "config_sha256": "cd" * 32, "git_sha": "deadbeef"}


@pytest.fixture(scope="session")
def xlsx_path(tmp_path_factory, abc, abc_metrics, abc_flagset):
    p = tmp_path_factory.mktemp("xlsx") / "trend_pack.xlsx"
    write_workbook(p, abc.company, abc_metrics, abc_flagset, abc.recon, abc.mapped, META)
    return p


@pytest.fixture(scope="session")
def rows(xlsx_path):
    """{sheet: list of row tuples} for every sheet except Data (read separately)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    pad = lambda r: list(r) + [None] * max(0, 12 - len(r))       # blank rows come back empty; keep widths uniform
    out = {ws.title: [pad(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets if ws.title != "Data"}
    wb.close()
    return out


@pytest.fixture(scope="session")
def data_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    out = [tuple(list(r) + [None] * (9 - len(r))) for r in wb["Data"].iter_rows(values_only=True)]   # trailing blanks are trimmed on read
    wb.close()
    return out


def find(rs, label, start=0, col=0):
    for i in range(start, len(rs)):
        if rs[i][col] == label:
            return i
    raise AssertionError(f"{label!r} not found")


def block(rs, title_prefix):
    """{label: row} for the table that follows the title row starting with `title_prefix`."""
    i = next(i for i, r in enumerate(rs) if isinstance(r[0], str) and r[0].startswith(title_prefix))
    out, j = {}, i + 2                                  # title, header, then rows until a blank row
    while j < len(rs) and rs[j][0] not in (None, ""):
        out[rs[j][0]] = rs[j]
        j += 1
    return out


def near(a, b, tol=0.005):
    assert isinstance(a, (int, float, Decimal)), a
    assert abs(float(a) - float(b)) <= tol, (a, b)


# ---------------------------------------------------------------- structure

def test_sheet_names_follow_plan_plus_data_accounts(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    assert wb.sheetnames == SHEETS == [
        "Cover", "Summary", "Revenue", "Margins", "Opex_by_function", "Opex_by_nature", "Opex_detail", "Windows",
        "Adjustments", "Flags", "Reconciliation", "Data", "Data_accounts", "Commentary", "Assumptions"]
    wb.close()


def test_values_only_no_formulas_no_macros_no_charts(xlsx_path):
    with zipfile.ZipFile(xlsx_path) as z:
        names = z.namelist()
        assert not [n for n in names if "vba" in n.lower() or "chart" in n.lower() or "drawing" in n.lower()]
        assert "macroEnabled" not in z.read("[Content_Types].xml").decode()
        for n in names:
            if n.startswith("xl/worksheets/sheet"):
                xml = z.read(n)
                assert b"<f>" not in xml and b"<f " not in xml, n
    assert xlsx_path.suffix == ".xlsx"


def test_cover(rows):
    cover = {r[0]: r[1] for r in rows["Cover"] if r[0] and len(r) > 1}
    text = " ".join(str(c) for r in rows["Cover"] for c in r if c)
    assert "DRAFT" in text
    assert cover["Company"] == "abc_inc" and cover["Status"] == "OK" and cover["As-of month"] == "2026-06"
    assert cover["Window lengths (months)"] == "12" and cover["Window end"] == "2026-06"
    assert cover["Input workbook SHA-256"] == META["input_sha256"] and cover["Config SHA-256"] == META["config_sha256"]
    assert cover["Git SHA"] == "deadbeef" and cover["First month"] == "2025-01"


def test_commentary_and_assumptions_sheets(rows):
    assert "not been generated" in " ".join(str(c) for r in rows["Commentary"] for c in r if c)
    assert len(rows["Assumptions"]) > 8


# ---------------------------------------------------------------- Summary: anchors, both bases, Dec 2025 both ways

def test_summary_latest_month_anchors(rows):
    t = block(rows["Summary"], "Latest month")
    assert t["Revenue"][2] == pytest.approx(195898.16) and t["Revenue"][3] == pytest.approx(195898.16)
    assert t["Revenue"][7] == pytest.approx(-12452.80)                 # MoM change reported
    near(t["Revenue"][8], -5.98, 0.005)                                  # MoM %
    near(t["Gross margin"][2], 70.15, 0.005)
    near(t["Total operating expenses as % of revenue"][2], 238.99, 0.005)
    assert t["Total operating expenses"][2] == pytest.approx(468181.80)
    assert t["Revenue"][5] == pytest.approx(208350.96)                  # prior month reported


def test_summary_header_labels_reported_and_adjusted_side_by_side(rows):
    i = next(i for i, r in enumerate(rows["Summary"]) if r[0] == "Metric")
    h = rows["Summary"][i]
    assert h[2] == "Reported" and h[3] == "Adjusted" and h[4].startswith("Delta")


def test_summary_ttm_window_anchors(rows):
    t = block(rows["Summary"], "12-month window ending 2026-06")
    assert t["Revenue"][2] == pytest.approx(2303975.51)
    assert t["Operating income"][2] == pytest.approx(-3089779.59) and t["Operating income"][3] == pytest.approx(-3999548.90)
    assert t["Operating income"][4] == pytest.approx(-909769.31)


def test_summary_dec_2025_shown_both_ways(rows):
    t = block(rows["Summary"], "Adjusted period 2025-12")
    oi = t["Operating income (month)"]
    assert oi[2] == pytest.approx(534551.41) and oi[3] == pytest.approx(-375217.90) and oi[4] == pytest.approx(-909769.31)
    near(t["Operating margin (month)"][2], 281.30, 0.005)
    near(t["Operating margin (month)"][3], -197.45, 0.005)
    w = t["Operating income (12-month window ending 2025-12)"]
    assert w[2] == pytest.approx(-3415098.97) and w[3] == pytest.approx(-4324868.28)
    ex = t["Excluded: Income Tax Benefit"]
    assert ex[2] == pytest.approx(-1244429.84) and ex[3] == 0 and ex[4] == pytest.approx(1244429.84)
    assert t["Excluded: Stock-Based Compensation Expense"][2] == pytest.approx(334660.53)


# ---------------------------------------------------------------- monthly sheets against the oracle

def matrix(rs, label, basis, view="value"):
    h = next(i for i, r in enumerate(rs) if r[0] == "Metric")
    periods = [p for p in rs[h][4:] if p]
    for r in rs[h + 1:]:
        if (r[0], r[1], r[2]) == (label, basis, view):
            return dict(zip(periods, r[4:4 + len(periods)]))
    raise AssertionError((label, basis, view))


def test_periods_run_oldest_to_newest(rows):
    h = next(r for r in rows["Revenue"] if r[0] == "Metric")
    ps = [p for p in h[4:] if p]
    assert ps == sorted(ps) and ps[0] == "2025-01" and ps[-1] == "2026-06" and len(ps) == 18


def test_revenue_sheet_equals_the_oracle_every_month(rows, abc):
    mo = oracle.monthly(abc.o)
    got = matrix(rows["Revenue"], "Revenue", "Reported")
    assert set(got) == set(mo)
    for p, v in got.items():
        near(v, mo[p]["rev"])
    adj = matrix(rows["Revenue"], "Revenue", "Adjusted")
    assert all(abs(float(adj[p]) - float(got[p])) < 1e-9 for p in got)          # revenue is not adjusted
    near(matrix(rows["Revenue"], "Revenue", "Reported", "mom_pct")["2026-06"], -5.98)


def test_margins_sheet_equals_the_oracle(rows, abc):
    mo = oracle.monthly(abc.o)
    gm = matrix(rows["Margins"], "Gross margin", "Reported")
    op = matrix(rows["Margins"], "Operating income", "Reported")
    zeroed = {(a.account, p) for a in abc.company.adjustments for p in a.periods}
    mo_adj = oracle.monthly(abc.o, zeroed)
    opa = matrix(rows["Margins"], "Operating income", "Adjusted")
    for p in mo:
        near(gm[p], oracle.pct(mo[p]["gp"], mo[p]["rev"]), 0.0001)
        near(op[p], mo[p]["opinc"])
        near(opa[p], mo_adj[p]["opinc"])
    near(op["2025-12"], 534551.41)
    near(opa["2025-12"], -375217.90)


def test_opex_by_nature_and_function_equal_the_oracle(rows, abc):
    mp = abc.dir / "mapping.yaml"
    nat = oracle.by_category(abc.o, mp, "nature")
    fn = oracle.by_category(abc.o, mp, "function")
    pb = matrix(rows["Opex_by_nature"], "Payroll & benefits", "Reported")
    sm = matrix(rows["Opex_by_function"], "Sales & Marketing", "Reported")
    for p in nat:
        near(pb[p], nat[p]["payroll_benefits"])
        near(sm[p], fn[p]["sales_marketing"])
    near(pb["2026-06"], 328377.23)
    pct = matrix(rows["Opex_by_nature"], "Payroll & benefits", "Reported", "pct_rev")
    near(pct["2026-06"], 167.63, 0.005)
    tot = matrix(rows["Opex_by_function"], "Total operating expenses", "Reported")
    for p in fn:                                                   # the groupings add up to the total, in the sheet
        s = sum(float(matrix(rows["Opex_by_function"], lab, "Reported")[p])
                for lab in [r[0] for r in rows["Opex_by_function"] if r[1] == "Reported" and r[2] == "value"
                            and r[0] != "Total operating expenses"])
        assert abs(s - float(tot[p])) < 0.01


def test_opex_detail_has_every_account(rows, abc):
    labels = {r[0] for r in rows["Opex_detail"] if r[1] == "Reported" and r[2] == "value"}
    opex = {l["label"] for l in abc.o["leaves"] if l["section"] == "opex"}
    assert labels == opex and len(labels) == 70


def test_windows_sheet_equals_the_oracle(rows, abc):
    mo = oracle.monthly(abc.o)
    rev = {p: mo[p]["rev"] for p in mo}
    rs = rows["Windows"]
    i = next(i for i, r in enumerate(rs) if isinstance(r[0], str) and r[0].startswith("12-month"))
    got = matrix(rs[i:], "Revenue", "Reported", "value")
    for p, v in got.items():
        want = oracle.window(rev, p, 12)
        if want is None:
            assert isinstance(v, str) and v.startswith("insufficient data"), (p, v)
        else:
            near(v, want)
    near(got["2026-06"], 2303975.51)
    assert got["2025-11"].startswith("insufficient data (missing: 2024-12")


# ---------------------------------------------------------------- tables

def test_adjustments_sheet(rows):
    h = find(rows["Adjustments"], "Account")
    body = [r for r in rows["Adjustments"][h + 1:] if r[0]]
    assert {(r[0], r[1]) for r in body} == {("Income Tax Benefit", "2025-12"), ("Stock-Based Compensation Expense", "2025-12")}
    itb = next(r for r in body if r[0] == "Income Tax Benefit")
    assert itb[2] == pytest.approx(-1244429.84) and itb[3] == 0 and itb[4] == pytest.approx(1244429.84) and itb[5] in (True, "yes")


def test_flags_sheet_lists_every_flag(rows, abc_flagset):
    rs = rows["Flags"]
    h = find(rs, "Flag id")
    body = []
    for r in rs[h + 1:]:
        if not r[0]:
            break
        body.append(r)
    assert [r[0] for r in body] == [f.id for f in abc_flagset.flags]
    assert {r[1] for r in body} == {f.type for f in abc_flagset.flags}
    spike = next(r for r in body if r[1] == "ONE_OFF_SPIKE" and r[5] == "Bonus")
    assert "continued_months=5" in spike[6] and "direction=up" in spike[6]
    head = " ".join(str(c) for r in rs[:h] for c in r if c)
    assert "NEGATIVE_VALUE" in head and "Bank Service Charges" in " ".join(str(c) for r in rs[h:] for c in r if c)


def test_reconciliation_sheet_shows_every_control_and_passes(rows, abc):
    rs = rows["Reconciliation"]
    h = find(rs, "Control")
    body = []
    for r in rs[h + 1:]:
        if not r[0]:
            break
        body.append(r)
    assert len(body) == len(abc.recon.results) == 180
    assert {r[6] for r in body} == {"OK"}
    assert {r[0] for r in body} == {f"C{i}" for i in range(1, 9)} | {"C9a", "C9b"}
    jun_inc = next(r for r in body if r[0] == "C1" and r[2] == "2026-06")
    assert jun_inc[3] == pytest.approx(195898.16) and jun_inc[4] == pytest.approx(195898.16) and jun_inc[5] == 0
    assert re.fullmatch(r"Income Statement!C\d+", jun_inc[7])


# ---------------------------------------------------------------- Data sheet (one row per fact) and Data_accounts

def test_data_sheet_has_one_row_per_fact(data_rows, abc_metrics):
    assert data_rows[0] == ("Fact id", "Basis", "Metric", "Dimension", "View", "Period", "Value", "Unit", "Note")
    body = data_rows[1:]
    assert len(body) == len(abc_metrics.facts) > 50_000
    assert [r[0] for r in body] == list(abc_metrics.facts)               # same ids, same order: deterministic


def test_every_data_metric_id_exists_and_columns_agree_with_the_id(data_rows, abc_metrics):
    ids = set()
    for fid, basis, metric, dim, view, period, value, unit, note in data_rows[1:]:
        assert fid in abc_metrics.facts
        p = parse_fact_id(fid)
        names = {"rep": "reported", "adj": "adjusted", "delta": "delta"}
        assert (basis, metric, dim or "", view or "", period) == (names[p["basis"]], p["metric"], p["dimension"], p["view"], p["period"])
        assert unit == unit_of(p)
        ids.add(fid)
    assert len(ids) == len(abc_metrics.facts)


def test_data_values_equal_the_metrics_and_n_a_rows_explain_themselves(data_rows, abc_metrics):
    n_na = 0
    for r in data_rows[1::37]:                                           # every 37th row
        v = abc_metrics.facts[r[0]]
        if isinstance(v, NA):
            assert r[6] in (None, "") and r[8], r
            n_na += 1
        else:
            near(r[6], float(v), 0.00006 if r[7] != "usd" else 0.005)
    assert n_na > 10
    anchors = {r[0]: r for r in data_rows[1:] if r[0] in ("rep.rev.total.2026-06", "rep.opinc.2025-12",
                                                          "adj.opinc.2025-12", "rep.rev.total.tw12.2026-06")}
    assert anchors["rep.rev.total.2026-06"][6] == pytest.approx(195898.16)
    assert anchors["rep.opinc.2025-12"][6] == pytest.approx(534551.41)
    assert anchors["adj.opinc.2025-12"][6] == pytest.approx(-375217.90)
    assert anchors["rep.rev.total.tw12.2026-06"][6] == pytest.approx(2303975.51)


def test_data_accounts_sheet_is_pivot_ready(xlsx_path, abc):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    rs = [r for r in wb["Data_accounts"].iter_rows(values_only=True)]
    wb.close()
    assert rs[0] == ("Basis", "Period", "Section", "Account", "Path", "Function", "Nature", "Amount", "Revenue", "Group")
    body = rs[1:]
    assert len(body) == len(abc.records) * 2
    tax = [r for r in body if r[3] == "Income Tax Benefit" and r[1] == "2025-12"]
    assert {r[0]: r[7] for r in tax} == {"reported": pytest.approx(-1244429.84), "adjusted": 0}
    opex_rep_jun = sum(r[7] for r in body if r[0] == "reported" and r[1] == "2026-06" and r[2] == "opex")
    near(opex_rep_jun, 468181.80)
    assert next(r for r in body if r[3] == "Rent")[5] == "g_and_a"


# ---------------------------------------------------------------- determinism

def test_two_builds_have_identical_cell_values(tmp_path, abc, abc_metrics, abc_flagset, rows, data_rows):
    p2 = tmp_path / "again.xlsx"
    write_workbook(p2, abc.company, abc_metrics, abc_flagset, abc.recon, abc.mapped, META)
    wb = openpyxl.load_workbook(p2, read_only=True)
    checked = 0
    for ws in wb.worksheets:
        if ws.title == "Data":
            got = [tuple(list(r) + [None] * (9 - len(r))) for r in ws.iter_rows(values_only=True)]
            assert got == data_rows
        else:
            got = [list(r) + [None] * max(0, 12 - len(r)) for r in ws.iter_rows(values_only=True)]
            assert got == rows[ws.title], ws.title
        checked += 1
    wb.close()
    assert checked == 15
